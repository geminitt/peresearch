"""Read one file into text blocks that remember where they came from (line, page or notebook cell).

Nothing here decides what may be read; `guard.allowed` does that before a path gets here.

Text is read in the encoding it was written in: a byte order mark names UTF-8, UTF-16 or UTF-32; valid UTF-8 is read
as is; otherwise the most likely legacy encoding (GBK, Big5, Shift-JIS, EUC-KR, cp1258…) is used when a language can
be recognized in the decoded text or, for a text of MIN_GUESS_BYTES or more, when the decoded text holds no character
from the rare ranges wrong guesses land in; the guess is named in the note. Otherwise (a few bytes are not enough) the
text is read as UTF-8 with replacement characters, and the note says so. Decoded legacy text is composed (NFC).
A PDF is limited by the text it holds, not its bytes (a book of scans is large and short), and lines in the top or
bottom MARGIN of a page that recur on most pages, digits ignored (running heads, page numbers), are dropped; text
elsewhere on the page is never dropped, however often it recurs. A notebook's code cell keeps its text outputs.
Office documents and web pages are converted to text as Microsoft's MarkItDown does: Word to Markdown (headings,
paragraphs, lists, tables; cut by its sections), PowerPoint a page per slide (title, text, tables, the speaker's
notes), Excel a tab-separated table per sheet (data, like a .csv), an HTML page to its main text in Markdown
(trafilatura) unless that holds little text (a template is code, read as it is).
Lines are numbered as editors number them (`lines`).
"""

import codecs
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from peresearch.files import special

MAX_BYTES = 20 * 2**20            # a text file larger than this is data, not notes
MAX_PDF_BYTES = 2**30             # a sanity bound only: a PDF is limited by its text
MAX_TEXT_CHARS = 20 * 2**20       # text read from one PDF
MAX_OUTPUT_CHARS = 2000           # text outputs kept per notebook cell for the index (read shows them whole)
# Top and bottom share of a PDF page where running heads and page numbers sit. Measured 2026-10-03 on 5 arXiv
# papers: heads end by 6%, page numbers start at 93.7%, body text reaches 7.7% from the top; at 10% a body line was
# dropped. Word's defaults agree: headers 0.5 in from the edge, body text 1 in (8.5% of an A4 page).
MARGIN = 0.075
MIN_GUESS_BYTES = 64              # below this, only a guess whose language is recognized is trusted
BOMS = [(codecs.BOM_UTF32_LE, "utf-32", "UTF-32"), (codecs.BOM_UTF32_BE, "utf-32", "UTF-32"),   # before UTF-16's
        (codecs.BOM_UTF8, "utf-8-sig", "UTF-8"), (codecs.BOM_UTF16_LE, "utf-16", "UTF-16"),
        (codecs.BOM_UTF16_BE, "utf-16", "UTF-16")]
_UNMAPPED = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\x7f]")   # a PDF glyph with no character: shown as U+FFFD
TEXT = {".md", ".markdown", ".txt", ".rst", ".org", ".tex", ".bib"}   # data files (csv, json, logs) are not notes
CODE = {".py", ".pyi", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".rs", ".go", ".java", ".kt", ".scala", ".c", ".h",
        ".cc", ".cpp", ".hpp", ".cs", ".rb", ".php", ".swift", ".m", ".r", ".jl", ".lua", ".sh", ".bash", ".zsh",
        ".fish", ".sql", ".toml", ".yaml", ".yml", ".ini", ".cfg", ".css", ".scss", ".cu"}
MARKDOWN = {".md", ".markdown"}
DATA = {".csv", ".tsv", ".json", ".jsonl", ".log"}   # results and logs: read as text by the tools, not ranked
OTHER = {".pdf": "pdf", ".ipynb": "notebook", ".docx": "docx", ".pptx": "pptx", ".xlsx": "spreadsheet",
         ".html": "html", ".htm": "html"}
CONVERTED = {"pdf", "notebook", "docx", "pptx", "spreadsheet", "html"}   # their text is not their bytes
PACKED = {"pdf", "docx", "pptx", "spreadsheet"}   # limited by their text (images make them large), not their bytes
# A kind's reader version: raise it when the reader changes what it reads, and the index reads that kind's files
# again (html 2: the main text of a page instead of its markup, 2026-10-05; docx 2: files holding XML comments read,
# pptx 2: charts read, 2026-10-06)
READERS = {"text": 1, "markdown": 1, "data": 1, "pdf": 1, "notebook": 1, "docx": 2, "pptx": 2, "spreadsheet": 1,
           "html": 2}
HTML_MIN_TEXT = 200          # characters of main text an HTML page needs to be read as a page rather than as code


@dataclass
class Block:
    text: str
    unit: str        # "line", "page" or "cell"
    start: int       # 1-based line, page or cell number where `text` begins


@dataclass
class Doc:
    path: Path
    status: str      # ok | needs_ocr | too_large | binary | unsupported | error
    blocks: list[Block] = field(default_factory=list)
    note: str = ""
    markdown: bool = False   # its text has Markdown headings to cut it by (a .md file, Word, a web page's text)


def kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in MARKDOWN:
        return "markdown"
    if ext in TEXT or ext in CODE:
        return "text"
    if ext in DATA:
        return "data"
    return OTHER.get(ext, "unsupported")


def read(path: Path, cut_outputs: bool = True) -> Doc:
    """A file's text in blocks. `cut_outputs` keeps MAX_OUTPUT_CHARS of a notebook cell's outputs, as the index and
    grep (whose two engines must agree with the index's extracted texts) see them; the read tool shows them whole."""
    path = Path(path)
    k = kind(path)
    if k == "unsupported":
        return Doc(path, "unsupported")
    odd = special(path)
    if odd is not None:                       # a named pipe would wait for a writer for ever
        return Doc(path, "unsupported", note=odd)
    try:
        size = path.stat().st_size
        if size > (MAX_PDF_BYTES if k in PACKED else MAX_BYTES):
            return Doc(path, "too_large", note=f"{size} bytes")
        if k == "pdf":
            return _pdf(path)
        if k in ("docx", "pptx", "spreadsheet"):
            return {"docx": _docx, "pptx": _pptx, "spreadsheet": _spreadsheet}[k](path)
        raw = path.read_bytes()
        bom = next(((c, label) for b, c, label in BOMS if raw.startswith(b)), None)
        if bom is None and b"\x00" in raw[:8192]:
            return Doc(path, "binary")
        text, note = decode(raw)
        if k == "notebook":
            return _notebook(path, text, cut_outputs)
        if k == "html":
            page = _html(text)
            if page is not None:
                return Doc(path, "ok", [Block(page, "line", 1)], note="main text of the page", markdown=True)
        return Doc(path, "ok", [Block(text, "line", 1)], note=note, markdown=k == "markdown")
    except Exception as e:  # a broken file is reported, never fatal
        return Doc(path, "error", note=f"{type(e).__name__}: {e}"[:200])


def lines(text: str) -> list[str]:
    """The lines of `text` as editors, git and ripgrep number them, the numbering every "L12" of peresearch uses: a
    line ends at "\n" (a "\r" before it belongs to the break). A form feed, U+2028 and the other characters
    str.splitlines() also breaks at stay inside their line."""
    out = text.split("\n")
    if out[-1] == "":
        out.pop()
    return [x[:-1] if x.endswith("\r") else x for x in out]


def decode(raw: bytes) -> tuple[str, str]:
    """(text, note) for a file's bytes; see the module's docstring."""
    for bom, codec, label in BOMS:
        if raw.startswith(bom):
            return raw.decode(codec), f"{label} (byte order mark)"
    try:
        return raw.decode("utf-8"), ""
    except UnicodeDecodeError:
        pass
    from charset_normalizer import from_bytes

    best = from_bytes(raw).best()
    # Measured 2026-10-03: wrong guesses (8–18 byte texts) all scored coherence 0 and produced rare characters
    # (CJK extension B, Arabic presentation forms); right Chinese and Vietnamese guesses scored above 0, but right
    # Japanese and Korean ones scored 0 too (charset-normalizer recognizes no language there), hence the second test.
    if best is not None and (best.coherence > 0 or (len(raw) >= MIN_GUESS_BYTES and _plausible(str(best)))):
        text = unicodedata.normalize("NFC", str(best))
        assert "\x00" not in text, f"{best.encoding} decoded NUL characters"
        return text, f"decoded as {best.encoding} (guessed)"
    return (raw.decode("utf-8", errors="replace"),
            "not UTF-8 and the encoding was not recognized: unreadable bytes replaced")


def _plausible(text: str) -> bool:
    """No character from the ranges a wrong guess lands in: planes 2 and above, private use, Arabic presentation
    forms."""
    return not any(ord(c) >= 0x20000 or 0xE000 <= ord(c) <= 0xF8FF or 0xFB50 <= ord(c) <= 0xFDFF
                   or 0xFE70 <= ord(c) <= 0xFEFC for c in text)


def _boilerplate(pages: list[list[str]]) -> set[str]:
    """Margin lines (running heads, page numbers) found on most pages, digits ignored. `pages` holds each page's
    margin lines; a document of fewer than 3 pages has too few to tell."""
    if len(pages) < 3:
        return set()
    seen: dict[str, int] = {}
    for lines in pages:
        for key in {_key(x) for x in lines}:
            seen[key] = seen.get(key, 0) + 1
    return {k for k, n in seen.items() if n >= max(3, len(pages) / 2)}


def _key(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip())


def _blocks(page) -> list[tuple[str, bool]]:
    """A page's text blocks in PyMuPDF's order, each with whether it lies in the top or bottom MARGIN of the page
    (joined, they are the page's plain text, measured equal on 138 of 141 pages of 5 arXiv papers; the 3 others
    differ by a line break around a math glyph)."""
    h = page.rect.height
    return [(_UNMAPPED.sub("\ufffd", text), y1 <= MARGIN * h or y0 >= (1 - MARGIN) * h)
            for x0, y0, x1, y1, text, _, kind in page.get_text("blocks") if kind == 0]


def read_pdf(body: bytes, name: str) -> Doc:
    """A PDF that arrived as bytes (from the web), read as a file would be; `name` is where it came from."""
    try:
        return _pdf(Path(name), body)
    except Exception as e:  # a broken file is reported, never fatal
        return Doc(Path(name), "error", note=f"{type(e).__name__}: {e}"[:200])


def _pdf(path: Path, body: bytes | None = None) -> Doc:
    import pymupdf

    # MuPDF prints its complaints about a damaged file on stdout — through the TUI's screen and the index report
    # (reproduced 2026-10-08): kept, counted in the note instead
    pymupdf.TOOLS.mupdf_display_errors(False)
    pymupdf.TOOLS.mupdf_display_warnings(False)
    pymupdf.TOOLS.reset_mupdf_warnings()
    doc = _pdf_read(pymupdf, path, body)
    problems = len([x for x in pymupdf.TOOLS.mupdf_warnings(reset=True).splitlines() if x.strip()])
    if problems:
        doc.note = (doc.note + "; " if doc.note else "") + f"MuPDF reported {problems} problem{'s' * (problems != 1)}"
    return doc


def _pdf_read(pymupdf, path: Path, body: bytes | None) -> Doc:
    with (pymupdf.open(stream=body, filetype="pdf") if body is not None else pymupdf.open(path)) as pdf:
        pages, chars = [], 0
        for i, page in enumerate(pdf):
            pages.append(_blocks(page))
            chars += sum(len(t) for t, _ in pages[-1])
            if chars > MAX_TEXT_CHARS:
                return Doc(path, "too_large", note=f"over {MAX_TEXT_CHARS} characters of text by page {i + 1}")
        n_images = sum(len(page.get_images()) for page in pdf)
    boilerplate = _boilerplate([[x for t, edge in blocks if edge for x in t.splitlines() if x.strip()]
                                for blocks in pages])
    texts = ["".join(t if not edge else "".join(x for x in t.splitlines(keepends=True) if _key(x) not in boilerplate)
                     for t, edge in blocks) for blocks in pages]
    blocks = [Block(t, "page", n) for n, t in enumerate(texts, start=1) if t.strip()]
    assert all(a.start < b.start for a, b in zip(blocks, blocks[1:])), "pages out of order"
    if not blocks:
        note = f"{len(pages)} pages, no text layer" + (f", {n_images} images" if n_images else "")
        return Doc(path, "needs_ocr", note=note)
    empty = len(pages) - len(blocks)
    return Doc(path, "ok", blocks, note=f"{empty} pages without text" if empty else "")


def _notebook(path: Path, text: str, cut_outputs: bool = True) -> Doc:
    cells = json.loads(text).get("cells", [])
    blocks = []
    for i, cell in enumerate(cells, start=1):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else src
        out = _outputs(cell, cut_outputs) if cell.get("cell_type") == "code" else ""
        if src.strip() and cell.get("cell_type") in ("markdown", "code"):
            blocks.append(Block(src + (f"\n\nOutput:\n{out}" if out else ""), "cell", i))
    return Doc(path, "ok", blocks)


def _outputs(cell: dict, cut: bool = True) -> str:
    """A code cell's text outputs (printed text, results, errors), cut at MAX_OUTPUT_CHARS when `cut`; images are left
    out."""
    parts = []
    for o in cell.get("outputs") or []:
        if not isinstance(o, dict):
            continue
        text = o.get("text")
        if text is None:
            text = (o.get("data") or {}).get("text/plain")
        if text is None and o.get("output_type") == "error":
            text = f"{o.get('ename', 'Error')}: {o.get('evalue', '')}"
        if text is not None:
            parts.append(("".join(text) if isinstance(text, list) else str(text)).rstrip("\n"))
    out = "\n".join(parts).strip()
    if cut and len(out) > MAX_OUTPUT_CHARS:
        out = out[:MAX_OUTPUT_CHARS] + f"\n… [{len(out) - MAX_OUTPUT_CHARS} more characters of output]"
    return out


def _limit(size: int) -> None:
    """A converted document is limited by its text, as a PDF is (MAX_TEXT_CHARS): a read error past it."""
    if size > MAX_TEXT_CHARS:
        raise ValueError(f"over {MAX_TEXT_CHARS} characters of text")


def _cell(text: str) -> str:
    """A table cell on one line, its bars escaped, for a Markdown table row."""
    return " ".join(text.split()).replace("|", "\\|")


def _table(rows: list[list[str]]) -> list[str]:
    """A Markdown table, the first row as its header (Word and PowerPoint give every row all its cells)."""
    if not rows:
        return []
    line = lambda r: "| " + " | ".join(_cell(c) for c in r) + " |"           # noqa: E731
    return [line(rows[0]), "| " + " | ".join(["---"] * len(rows[0])) + " |", *map(line, rows[1:])]


def _docx(path: Path) -> Doc:
    """A Word document as Markdown, its paragraphs and tables in their order: "Heading n" and "Title" styles become
    headings, list styles and numbered paragraphs list items, tables Markdown tables."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    items: list[tuple[bool, list[str]]] = []                  # (a list item, its lines): items of a list stay together
    for child in d.element.body.iterchildren():
        if not isinstance(child.tag, str):                    # an XML comment or processing instruction: no text
            continue
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            rows = [[c.text for c in r.cells] for r in Table(child, d).rows]
            items.append((False, _table(rows)))
        elif tag == "p":
            p = Paragraph(child, d)
            text = " ".join(p.text.split("\n")).strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            level = re.match(r"Heading (\d)$", style)
            if style == "Title" or level:
                items.append((False, ["#" * (int(level.group(1)) if level else 1) + " " + text]))
            elif "List" in style or child.pPr is not None and child.pPr.numPr is not None:
                items.append((True, [("1. " if "Number" in style else "- ") + text]))
            else:
                items.append((False, [text]))
    out: list[str] = []
    for i, (listed, lines_) in enumerate(items):
        if i and not (listed and items[i - 1][0]):
            out.append("")
        out += lines_
    text = "\n".join(out)
    _limit(len(text))
    return Doc(path, "ok", [Block(text, "line", 1)] if out else [], note="converted from Word", markdown=True)


def _chart(chart) -> list[str]:
    """A chart as MarkItDown gives it: "Chart: <title>", then a table of its values, a row per category and a column
    per series. A chart whose values cannot be read as a table keeps its title alone: a scatter or bubble chart has
    no categories, and python-pptx raises NotImplementedError on kinds it does not know (a 3-D area chart) — any
    error, as MarkItDown catches, since one chart must not cost the deck its text."""
    try:
        title = (chart.chart_title.text_frame.text.strip()
                 if chart.has_title and chart.chart_title.has_text_frame else "")
    except Exception:
        title = ""
    out = [f"Chart: {title}" if title else "Chart"]
    try:
        categories = [str(c) for c in chart.plots[0].categories]
        series = list(chart.series)
        values = [list(s.values) for s in series]
        if categories and series:
            out += _table([["Category", *(s.name for s in series)],
                           *([c, *("" if v[i] is None else str(v[i]) for v in values)]
                             for i, c in enumerate(categories))])
    except Exception:
        pass
    return out


def _pptx(path: Path) -> Doc:
    """A PowerPoint deck, a page per slide that holds text: its title, the text of its shapes (groups included),
    tables and charts in their order, then the speaker's notes."""
    import pptx

    def shapes(group) -> list[str]:
        out = []
        for shape in group:
            if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                out += shapes(shape.shapes)
            elif shape.has_text_frame and shape.text_frame.text.strip() and shape != title:
                out.append(shape.text_frame.text.strip())
            elif shape.has_table:
                out += _table([[c.text for c in r.cells] for r in shape.table.rows])
            elif shape.has_chart:
                out += _chart(shape.chart)
        return out

    deck, blocks, total = pptx.Presentation(str(path)), [], 0
    for n, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title
        parts = ([title.text_frame.text.strip()] if title is not None and title.text_frame.text.strip() else [])
        parts += shapes(slide.shapes)
        text = "\n".join(parts)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            text += ("\n\n" if text else "") + f"Notes: {notes}"
        if text:
            total += len(text)
            _limit(total)
            blocks.append(Block(text, "page", n))
    return Doc(path, "ok", blocks, note="converted from PowerPoint")


def _spreadsheet(path: Path) -> Doc:
    """An Excel workbook, a tab-separated table per sheet that holds values (cached formula results, not formulas),
    each under "## <sheet>"; empty rows and trailing empty cells left out."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        out, size = [], 0
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else " ".join(str(v).split()) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    rows.append("\t".join(cells))
                    size += len(rows[-1])
                    _limit(size)
            if rows:
                out += ([""] if out else []) + [f"## {ws.title}"] + rows
    finally:
        wb.close()
    return Doc(path, "ok", [Block("\n".join(out), "line", 1)] if out else [], note="converted from Excel")


def _html(text: str) -> str | None:
    """A web page's main text in Markdown, or None when it holds less than HTML_MIN_TEXT characters of it (a
    template, a fragment: read as code)."""
    import trafilatura

    page = trafilatura.extract(text, output_format="markdown", include_tables=True, include_links=False,
                               include_comments=False)
    return page if page and len(page.strip()) >= HTML_MIN_TEXT else None
